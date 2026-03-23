"""
测试 xlsx 导入解析逻辑
验证字段映射是否正确，不写入数据库
"""
import sys
import io
import json

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from openpyxl import load_workbook

# ── 与导出保持对应的反向映射 ──────────────────────────────────
_XLSX_TYPE_REVERSE = {
    "单选题": "单选",
    "多选题": "多选",
    "判断题": "是非",
    "简答题": "简答",
    "计算题": "简答>计算",
    "论述题": "简答>论述",
    "材料分析题": "简答>材料分析",
}

def _diff_map(n):
    """难度数字 → 字符串"""
    try:
        n = int(float(n))
    except (TypeError, ValueError):
        return "medium"
    if n <= 2:
        return "easy"
    if n == 3:
        return "medium"
    return "hard"

def _parse_tags(tag_str):
    """'#知识点1#知识点2' → (knowledge_point, tags_str)"""
    if not tag_str:
        return None, None
    parts = [p.strip() for p in tag_str.split('#') if p.strip()]
    if not parts:
        return None, None
    kp = parts[0]
    tags = ','.join(parts[1:]) if len(parts) > 1 else None
    return kp, tags

def _normalize_answer(q_type, raw_answer):
    """
    根据题型规范化答案，返回 (answer, reference_answer)
    - 单选: answer='A', reference_answer=''
    - 多选: 'A,B' / 'A;B' / 'AB' → answer='AB', reference_answer=''
    - 判断: 'true'/'正确'→'正确', 'false'/'错误'→'错误'
    - 简答: answer='', reference_answer=文本
    """
    raw = (raw_answer or '').strip()
    if q_type == '单选':
        return raw.upper() if raw else '', ''
    if q_type == '多选':
        # 'A,B' / 'A;B' / 'A B' → 'AB'
        import re
        letters = re.findall(r'[A-Oa-o]', raw)
        return ''.join(l.upper() for l in letters), ''
    if q_type == '是非':
        if raw.lower() in ('true', '正确', '对', '是'):
            return '正确', ''
        if raw.lower() in ('false', '错误', '错', '否'):
            return '错误', ''
        return raw, ''
    # 简答类
    return '', raw

def parse_xlsx_questions(file_path):
    """
    解析 xlsx 题库文件，返回题目 dict 列表。
    自动检测标题行（row3 或 row1）。
    """
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    # 自动检测标题行：找包含 '题干' 的行
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), 1):
        row_str = ' '.join(str(c or '') for c in row)
        if '题干' in row_str:
            header_row = i
            break

    if header_row is None:
        raise ValueError("未找到标题行（含'题干'的行），请检查文件格式")

    data_start = header_row + 1
    questions = []
    errors = []

    for row_idx, row in enumerate(
        ws.iter_rows(min_row=data_start, values_only=True), data_start
    ):
        # 跳过空行（题干为空）
        stem = str(row[7] or '').strip() if len(row) > 7 else ''
        if not stem:
            continue

        raw_type = str(row[1] or '').strip()
        q_type = _XLSX_TYPE_REVERSE.get(raw_type)
        if not q_type:
            errors.append(f"Row {row_idx}: 未知题型 '{raw_type}'，已跳过")
            continue

        difficulty = _diff_map(row[3])
        tag_str = str(row[5] or '').strip()
        kp, tags = _parse_tags(tag_str)
        raw_answer = str(row[8] or '').strip() if len(row) > 8 else ''
        explanation = str(row[9] or '').strip() if len(row) > 9 else ''

        answer, reference_answer = _normalize_answer(q_type, raw_answer)

        # 选项：col 11-25 (index 10-24)
        options = []
        for ci in range(10, 25):
            val = row[ci] if ci < len(row) else None
            if val is not None:
                options.append(str(val).strip())

        questions.append({
            'type': q_type,
            'content': stem,
            'options': options,
            'answer': answer,
            'reference_answer': reference_answer,
            'explanation': explanation,
            'knowledge_point': kp,
            'tags': tags,
            'difficulty': difficulty,
        })

    return questions, errors


# ── 运行测试 ──────────────────────────────────────────────────
if __name__ == '__main__':
    file = '我的题库_管理研究方法_图谱题库_2026-03-21.xlsx'
    print(f'解析文件: {file}')
    print('=' * 60)

    questions, errors = parse_xlsx_questions(file)

    if errors:
        print('[警告]')
        for e in errors:
            print(' ', e)
        print()

    print(f'解析成功: {len(questions)} 题')
    print()

    # 按题型统计
    from collections import Counter
    types = Counter(q['type'] for q in questions)
    print('题型分布:', dict(types))
    print()

    # 展示各题型首题
    shown = set()
    for q in questions:
        t = q['type']
        if t not in shown:
            shown.add(t)
            print(f'--- [{t}] 示例 ---')
            print(f'  题干: {q["content"][:60]}...' if len(q["content"]) > 60 else f'  题干: {q["content"]}')
            print(f'  选项: {q["options"]}')
            print(f'  答案: {repr(q["answer"])}')
            print(f'  参考答案: {repr(q["reference_answer"])}')
            print(f'  解析: {repr(q["explanation"])}')
            print(f'  难度: {q["difficulty"]}')
            print(f'  知识点: {q["knowledge_point"]}')
            print(f'  标签: {q["tags"]}')
            print()
